#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build.py — 產生「北歐接力精確體系」靜態網站

來源：
  原始資料/一方塊開叫.xlsx, 一梅花開叫.xlsx, 高花開叫.xlsx   -> 網站/叫序資料/*.html
  原始資料/3基本準則/第三章.md
  原始資料/4一方塊開叫/一方塊開叫.md
  原始資料/5高花開叫/高花開叫.md
  原始資料/7一梅花開叫/一梅花開叫.md                          -> 網站/原始書籍檔案/*.html

重新執行本腳本即可依來源檔案的最新內容重新產生整個 網站/ 目錄。
"""
import os
import re
import html
import shutil

import openpyxl
import markdown as md

ROOT = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(ROOT, "原始資料")
OUT = os.path.join(ROOT, "網站")

SUIT_RED = set("♥♦")
SUIT_BLACK = set("♠♣")


# --------------------------------------------------------------------------
# 共用版型
# --------------------------------------------------------------------------

NAV_ITEMS = [
    ("叫序資料", [
        ("一方塊開叫", "叫序資料/一方塊開叫.html"),
        ("一梅花開叫", "叫序資料/一梅花開叫.html"),
        ("高花開叫", "叫序資料/高花開叫.html"),
    ]),
    ("原始書籍檔案", [
        ("第三章", "原始書籍檔案/第三章.html"),
        ("一方塊開叫", "原始書籍檔案/一方塊開叫.html"),
        ("高花開叫", "原始書籍檔案/高花開叫.html"),
        ("一梅花開叫", "原始書籍檔案/一梅花開叫.html"),
    ]),
]


def suit_colorize(text):
    """把文字中的 ♥♦ 轉紅、♠♣ 轉黑（沿用速查頁的 .r/.b 慣例），並保留換行。"""
    out = []
    for ch in text:
        if ch in SUIT_RED:
            out.append('<span class="s r">%s</span>' % ch)
        elif ch in SUIT_BLACK:
            out.append('<span class="s b">%s</span>' % ch)
        elif ch == "\n":
            out.append("<br>")
        else:
            out.append(html.escape(ch))
    return "".join(out)


def render_nav(root_prefix, current_href=None):
    parts = []
    parts.append('<nav class="site-nav">')
    parts.append('<div class="nav-inner">')
    parts.append('<a class="nav-brand" href="%sindex.html">北歐接力精確體系</a>' % root_prefix)
    parts.append('<button class="nav-toggle" aria-label="開啟選單" type="button">'
                  '<svg viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2">'
                  '<path d="M3 6h18M3 12h18M3 18h18"/></svg></button>')
    parts.append('<ul class="nav-links" role="menubar">')
    for label, children in NAV_ITEMS:
        parts.append('<li class="nav-item has-menu" role="none">')
        parts.append('<button class="nav-link" type="button">%s <span class="car">▾</span></button>' % label)
        parts.append('<ul class="nav-menu">')
        for sub_label, sub_href in children:
            active = " is-active" if sub_href == current_href else ""
            parts.append('<li><a class="%s" href="%s%s">%s</a></li>' %
                          (active.strip() or "", root_prefix + sub_href, "", sub_label))
        parts.append('</ul></li>')
    parts.append('</ul>')
    parts.append('</div></nav>')
    return "".join(parts)


PAGE_TEMPLATE = """<!doctype html>
<html lang="zh-Hant">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{title}</title>
<meta name="description" content="{description}">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=JetBrains+Mono:wght@500;700&family=Noto+Sans+TC:wght@400;500;700&family=Noto+Serif+TC:wght@600;700&display=swap">
<link rel="stylesheet" href="{root}assets/style.css">
</head>
<body>
{nav}
<div class="wrap">
{body}
</div>
<script src="{root}assets/site.js"></script>
</body>
</html>
"""


def render_page(title, description, root_prefix, current_href, body_html):
    return PAGE_TEMPLATE.format(
        title=html.escape(title),
        description=html.escape(description),
        root=root_prefix,
        nav=render_nav(root_prefix, current_href),
        body=body_html,
    )


def page_head(eyebrow, h1, meta_lines, lede_html, crumb_html=""):
    meta = "<br>".join(meta_lines)
    return """
{crumb}
<header class="page-head">
  <div class="mast">
    <div>
      <div class="eyebrow">{eyebrow}</div>
      <h1>{h1}</h1>
    </div>
    <div class="mast-meta">{meta}</div>
  </div>
  <p class="lede">{lede}</p>
</header>
""".format(crumb=crumb_html, eyebrow=eyebrow, h1=h1, meta=meta, lede=lede_html)


def write_file(path, content):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)
    print("wrote", os.path.relpath(path, ROOT))


# --------------------------------------------------------------------------
# xlsx -> HTML（叫序資料）
#
# 每個工作表先被 analyze_sheet() 拆成「前導說明列／表頭／資料列／註腳」，
# 再依表頭形狀分派給對應的語意化 renderer（沿用網站既有的排版元件，
# 不再輸出 Excel 原始底色）。資料列中形如「→ 表4「A分支」」「見表 9」的
# 文字，會被 linkify_refs() 轉成可點擊的分頁跳轉連結。
# --------------------------------------------------------------------------

HEADER_NAVY = "1F3864"


def is_header_style(cell):
    """判斷這個儲存格是否具備原始表格「表頭列」的樣式（粗體＋深藍底），
    用來定位表頭列，以及「說明」表中重複出現的小節分隔列。"""
    if not (cell.font and cell.font.bold):
        return False
    fill = cell.fill
    if not (fill and fill.fill_type == "solid" and fill.fgColor):
        return False
    fg = fill.fgColor
    if fg.type != "rgb" or not fg.rgb:
        return False
    return fg.rgb[-6:].upper() == HEADER_NAVY


def cell_text(ws, r, c):
    v = ws.cell(row=r, column=c).value
    if v is None:
        return ""
    v = str(v)
    if v.startswith("'"):
        # Excel「強制文字」前綴的殘留單引號，不是內容的一部分
        v = v[1:]
    return v


def analyze_sheet(ws):
    max_col = ws.max_column
    max_row = ws.max_row

    header_row = None
    for r in range(1, max_row + 1):
        if is_header_style(ws.cell(row=r, column=1)):
            header_row = r
            break
    if header_row is None:
        header_row = 1

    leading = [cell_text(ws, r, 1).strip() for r in range(1, header_row)]
    leading = [t for t in leading if t]

    header = [cell_text(ws, header_row, c).strip() for c in range(1, max_col + 1)]
    while header and header[-1] == "":
        header.pop()

    rows = []
    for r in range(header_row + 1, max_row + 1):
        vals = [cell_text(ws, r, c) for c in range(1, max_col + 1)]
        if not any(v.strip() for v in vals):
            continue
        if is_header_style(ws.cell(row=r, column=1)):
            rows.append(("SUB", vals[0].strip()))
            continue
        rows.append(("DATA", vals))

    footnotes = []
    while rows and rows[-1][0] == "DATA":
        vals = rows[-1][1]
        nonempty = [i for i, v in enumerate(vals) if v.strip()]
        if nonempty == [0]:
            footnotes.insert(0, vals[0].strip())
            rows.pop()
        else:
            break

    return dict(leading=leading, header=header, rows=rows, footnotes=footnotes, max_col=max_col)


BID_TOKEN_RE = re.compile(r'(\d)(<span class="s [rb]">[♠♥♦♣]</span>|NT)')


def bidify(html_text):
    return BID_TOKEN_RE.sub(lambda m: '<span class="bid">%s</span>' % m.group(0), html_text)


BRANCH_CODE_RE = re.compile(r"^([A-Za-z])(?:[\s　]|$)")


def extract_branch_code(raw):
    if not raw:
        return None
    m = BRANCH_CODE_RE.match(raw.strip())
    return m.group(1) if m else None


REF_RE = re.compile(r"表\s*(\d+(?:[、,]\s*\d+)*)(?:「([^」]*)」)?")


def linkify_refs(raw, num_to_index, merged_indices, current_code):
    """把「→ 表4「A分支」」「見表 9」「見表 3、4」這類文字轉成可點擊的分頁跳轉連結。"""

    def repl(m):
        nums_str, label = m.group(1), m.group(2)
        nums = re.split(r"[、,]\s*", nums_str)
        links = []
        for ns in nums:
            idx = num_to_index.get(int(ns))
            if idx is None:
                links.append(ns)
                continue
            dom = "sheet-%d" % (idx + 1)
            attrs = 'class="jump-link" data-sheet="%s"' % dom
            if idx in merged_indices and current_code:
                attrs += ' data-anchor="branch-%d-%s"' % (idx + 1, current_code)
            links.append("<a %s>%s</a>" % (attrs, ns))
        suffix = ("「%s」" % label) if label else ""
        return "表" + "、".join(links) + suffix

    return REF_RE.sub(repl, raw)


def render_cell(raw, num_to_index, merged_indices, current_code=None):
    """統一的儲存格內容渲染：跳轉連結 → 花色上色／逃逸 → 叫品 mono 樣式。"""
    if raw is None:
        raw = ""
    linked = linkify_refs(raw, num_to_index, merged_indices, current_code)
    parts = re.split(r"(<[^>]+>)", linked)
    out = []
    for part in parts:
        if part.startswith("<"):
            out.append(part)
            continue
        seg = suit_colorize(part)
        seg = bidify(seg)
        out.append(seg)
    return "".join(out)


def render_lead_block(leading, num_to_index, merged_indices):
    if not leading:
        return ""
    parts = ['<div class="lead-block">']
    parts.append('<div class="lead-title">%s</div>' %
                  render_cell(leading[0], num_to_index, merged_indices))
    if len(leading) > 1:
        parts.append('<div class="lead-context">%s</div>' %
                      render_cell(leading[1], num_to_index, merged_indices))
    for extra in leading[2:]:
        parts.append('<div class="lead-def">%s</div>' %
                      render_cell(extra, num_to_index, merged_indices))
    parts.append("</div>")
    return "".join(parts)


def render_footnotes(footnotes, num_to_index, merged_indices):
    if not footnotes:
        return ""
    lines = [render_cell(f, num_to_index, merged_indices) for f in footnotes]
    return '<div class="trap">%s</div>' % "<br>".join(lines)


def render_ladder(analysis, num_to_index, merged_indices):
    header = analysis["header"]
    n = len(header)
    code_col = bool(header) and header[0] == "代號"

    parts = [render_lead_block(analysis["leading"], num_to_index, merged_indices)]
    ths = "".join("<th>%s</th>" % html.escape(h) for h in header)
    trs = []
    for kind, payload in analysis["rows"]:
        if kind == "SUB":
            trs.append('<tr class="subhead"><td colspan="%d">%s</td></tr>' %
                        (n, render_cell(payload, num_to_index, merged_indices)))
            continue
        vals = (payload + [""] * n)[:n]
        current_code = extract_branch_code(vals[0])
        tds = []
        for ci, raw in enumerate(vals):
            if ci == 0 and code_col:
                content = ('<span class="code-badge">%s</span>' % html.escape(raw.strip())
                            if raw.strip() else "&nbsp;")
            elif ci == 0 and current_code and len(raw.strip()) > 1:
                rest = raw.strip()[1:].strip()
                content = '<span class="code-badge">%s</span> %s' % (
                    html.escape(current_code), render_cell(rest, num_to_index, merged_indices, current_code))
            else:
                content = render_cell(raw, num_to_index, merged_indices, current_code)
            if not content:
                content = "&nbsp;"
            cls = ' class="cell-muted"' if ci == n - 1 else ""
            tds.append("<td%s>%s</td>" % (cls, content))
        trs.append("<tr>%s</tr>" % "".join(tds))

    parts.append('<div class="table-scroll"><table class="bidtbl"><thead><tr>%s</tr></thead>'
                  '<tbody>%s</tbody></table></div>' % (ths, "".join(trs)))
    parts.append(render_footnotes(analysis["footnotes"], num_to_index, merged_indices))
    return "".join(parts)


def render_merged_branch(analysis, sheet_index, num_to_index, merged_indices):
    header = analysis["header"]
    inner_headers = header[2:]

    groups = []
    for kind, payload in analysis["rows"]:
        if kind == "SUB":
            continue
        vals = (payload + [""] * len(header))[:len(header)]
        code = vals[0].strip()
        context = vals[1]
        rest = vals[2:]
        if not groups or groups[-1][0] != code:
            groups.append([code, context, [rest]])
        else:
            groups[-1][2].append(rest)

    parts = [render_lead_block(analysis["leading"], num_to_index, merged_indices)]
    for code, context, rows in groups:
        anchor_id = "branch-%d-%s" % (sheet_index + 1, code)
        parts.append('<div class="branch-group" id="%s">' % anchor_id)
        parts.append('<div class="branch-head"><span class="code-badge">%s</span><span class="ctx">%s</span></div>' %
                      (html.escape(code), render_cell(context, num_to_index, merged_indices)))
        ths = "".join("<th>%s</th>" % html.escape(h) for h in inner_headers)
        trs = []
        for rvals in rows:
            tds = []
            for ci, raw in enumerate(rvals):
                content = render_cell(raw, num_to_index, merged_indices, code) or "&nbsp;"
                cls = ' class="cell-muted"' if ci == len(rvals) - 1 else ""
                tds.append("<td%s>%s</td>" % (cls, content))
            trs.append("<tr>%s</tr>" % "".join(tds))
        parts.append('<div class="table-scroll"><table class="bidtbl"><thead><tr>%s</tr></thead>'
                      '<tbody>%s</tbody></table></div>' % (ths, "".join(trs)))
        parts.append("</div>")

    parts.append(render_footnotes(analysis["footnotes"], num_to_index, merged_indices))
    return "".join(parts)


def render_info(analysis, num_to_index, merged_indices):
    parts = [render_lead_block(analysis["leading"], num_to_index, merged_indices)]
    parts.append('<table class="deftable"><tbody>')
    for kind, payload in analysis["rows"]:
        if kind == "SUB":
            parts.append('</tbody></table><div class="deftable-sub">%s</div><table class="deftable"><tbody>' %
                          html.escape(payload))
            continue
        vals = (payload + ["", ""])[:2]
        label = render_cell(vals[0], num_to_index, merged_indices)
        desc = render_cell(vals[1], num_to_index, merged_indices)
        parts.append("<tr><td>%s</td><td>%s</td></tr>" % (label, desc))
    parts.append("</tbody></table>")
    parts.append(render_footnotes(analysis["footnotes"], num_to_index, merged_indices))
    return "".join(parts)


def render_hands(analysis, num_to_index, merged_indices):
    parts = [render_lead_block(analysis["leading"], num_to_index, merged_indices)]
    groups = []
    for kind, payload in analysis["rows"]:
        if kind == "SUB":
            continue
        vals = (payload + [""] * 10)[:10]
        exno = vals[0].strip()
        if not groups or groups[-1][0] != exno:
            groups.append([exno, []])
        groups[-1][1].append(vals)

    for exno, rows in groups:
        first = rows[0]
        source, situation = first[1], first[2]
        contract = next((r[8] for r in rows if r[8].strip()), "")
        comment = next((r[9] for r in rows if r[9].strip()), "")
        parts.append('<div class="example-card">')
        parts.append('<div class="example-head"><b>%s</b><span>%s</span><span>%s</span></div>' %
                      (html.escape(exno), render_cell(source, num_to_index, merged_indices),
                       render_cell(situation, num_to_index, merged_indices)))
        parts.append('<div class="example-body"><div class="table-scroll"><table class="handtbl"><tbody>')
        for r in rows:
            seat, sp, he, di, cl = r[3], r[4], r[5], r[6], r[7]
            parts.append(
                '<tr><td class="seat">%s</td>'
                '<td class="suit"><span class="s b">♠</span>%s</td>'
                '<td class="suit"><span class="s r">♥</span>%s</td>'
                '<td class="suit"><span class="s r">♦</span>%s</td>'
                '<td class="suit"><span class="s b">♣</span>%s</td></tr>' %
                (render_cell(seat, num_to_index, merged_indices),
                 render_cell(sp, num_to_index, merged_indices),
                 render_cell(he, num_to_index, merged_indices),
                 render_cell(di, num_to_index, merged_indices),
                 render_cell(cl, num_to_index, merged_indices)))
        parts.append("</tbody></table></div>")
        if contract.strip() or comment.strip():
            parts.append('<div class="example-result">')
            if contract.strip():
                parts.append('<span class="bid-lead">%s</span>' % render_cell(contract, num_to_index, merged_indices))
            if comment.strip():
                parts.append(render_cell(comment, num_to_index, merged_indices))
            parts.append("</div>")
        parts.append("</div></div>")

    parts.append(render_footnotes(analysis["footnotes"], num_to_index, merged_indices))
    return "".join(parts)


def render_auction(analysis, num_to_index, merged_indices):
    parts = [render_lead_block(analysis["leading"], num_to_index, merged_indices)]
    groups = []
    for kind, payload in analysis["rows"]:
        if kind == "SUB":
            continue
        vals = (payload + [""] * 6)[:6]
        exno = vals[0].strip()
        # 只有每個例子的第一列才有「例號」；後續列的例號欄位是空的，仍屬於同一組
        if exno and (not groups or groups[-1][0] != exno):
            groups.append([exno, vals[1], []])
        if not groups:
            groups.append(["", vals[1], []])
        groups[-1][2].append(vals[2:])

    for exno, situation, rows in groups:
        parts.append('<div class="example-card">')
        parts.append('<div class="example-head"><b>%s</b><span>%s</span></div>' %
                      (html.escape(exno), render_cell(situation, num_to_index, merged_indices)))
        parts.append('<div class="example-body"><div class="table-scroll"><table class="auctiontbl"><thead>'
                      '<tr><th>序</th><th>叫者</th><th>叫品</th><th>說明</th></tr></thead><tbody>')
        for seq, who, bid, note in rows:
            note_html = render_cell(note, num_to_index, merged_indices) if note.strip() else ""
            parts.append(
                '<tr><td class="seq">%s</td><td class="who">%s</td>'
                '<td>%s</td><td>%s</td></tr>' %
                (html.escape(seq.strip()), render_cell(who, num_to_index, merged_indices),
                 render_cell(bid, num_to_index, merged_indices), note_html))
        parts.append("</tbody></table></div></div></div>")

    parts.append(render_footnotes(analysis["footnotes"], num_to_index, merged_indices))
    return "".join(parts)


def render_sheet_panel(analysis, sheet_index, num_to_index, merged_indices):
    header = tuple(analysis["header"])
    if header[:2] == ("項目", "內容"):
        return render_info(analysis, num_to_index, merged_indices)
    if header[:2] == ("分支", "序列／前提"):
        return render_merged_branch(analysis, sheet_index, num_to_index, merged_indices)
    if header[:1] == ("例號",) and "方位" in header:
        return render_hands(analysis, num_to_index, merged_indices)
    if header[:1] == ("例號",) and "叫者" in header:
        return render_auction(analysis, num_to_index, merged_indices)
    return render_ladder(analysis, num_to_index, merged_indices)


def build_xlsx_page(xlsx_name, title, eyebrow, lede):
    path = os.path.join(SRC, xlsx_name)
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = wb.worksheets

    analyses = [analyze_sheet(ws) for ws in sheets]
    num_to_index = {}
    for i, ws in enumerate(sheets):
        m = re.match(r"^(\d+)\.", ws.title)
        if m:
            num_to_index[int(m.group(1))] = i
    merged_indices = {i for i, a in enumerate(analyses) if tuple(a["header"][:2]) == ("分支", "序列／前提")}

    tabs = []
    panels = []
    for i, ws in enumerate(sheets):
        sid = "sheet-%d" % (i + 1)
        active = " is-active" if i == 0 else ""
        tabs.append('<button class="sheet-tab%s" data-target="%s">%s</button>' %
                     (active, sid, html.escape(ws.title)))
        content = render_sheet_panel(analyses[i], i, num_to_index, merged_indices)
        panels.append('<div class="sheet-panel%s" id="%s">%s</div>' % (active, sid, content))

    body = page_head(
        eyebrow, title,
        ["%d 個工作表" % len(sheets), "來源：%s" % xlsx_name],
        lede,
        crumb_html='<div class="crumb"><a href="../index.html">首頁</a> ／ 叫序資料 ／ %s</div>' % title,
    )
    body += '<section>'
    body += '<div class="sheet-tabs" data-sheet-tabs="#%s-panels">%s</div>' % (title, "".join(tabs))
    body += '<div id="%s-panels">%s</div>' % (title, "".join(panels))
    body += '</section>'
    body += '<footer class="site-footer">資料整理自《北歐接力精確體系》，原始表格見「原始書籍檔案」分頁。</footer>'

    html_out = render_page(
        title="%s ｜ 叫序資料 ｜ 北歐接力精確體系" % title,
        description="%s 的完整叫序資料表" % title,
        root_prefix="../",
        current_href="叫序資料/%s.html" % title,
        body_html=body,
    )
    write_file(os.path.join(OUT, "叫序資料", "%s.html" % title), html_out)


# --------------------------------------------------------------------------
# markdown -> HTML（原始書籍檔案）
# --------------------------------------------------------------------------

FRONTMATTER_RE = re.compile(r"^---\s*\n.*?\n---\s*\n", re.S)


def strip_frontmatter(text):
    return FRONTMATTER_RE.sub("", text, count=1)


def md_to_html(md_path, image_dir_href):
    with open(md_path, "r", encoding="utf-8") as f:
        text = f.read()
    text = strip_frontmatter(text)

    md_conv = md.Markdown(extensions=["tables", "sane_lists"])
    body = md_conv.convert(text)

    # 圖片路徑改指到 assets/images/<章節>/
    def fix_img(m):
        pre, src, post = m.group(1), m.group(2), m.group(3)
        if src.startswith("http://") or src.startswith("https://"):
            return m.group(0)
        fname = os.path.basename(src)
        return '<figure><img%ssrc="%s%s"%s></figure>' % (pre, image_dir_href, fname, post)

    body = re.sub(r'<img([^>]*?)src="([^"]+)"([^>]*?)>', fix_img, body)
    # markdown 會把單獨成行的圖片包在 <p> 裡，figure 是區塊元素，拆掉外層 <p> 避免非法巢狀
    body = re.sub(r"<p>\s*(<figure>.*?</figure>)\s*</p>", r"\1", body, flags=re.S)

    # 幫表格加水平捲動包裹（RWD）
    body = body.replace("<table>", '<div class="table-scroll"><table>')
    body = body.replace("</table>", "</table></div>")

    body = colorize_suits_html(body)

    return body


def colorize_suits_html(html_body):
    """花色符號著色，只處理標籤之間的純文字節點，避免污染 alt= 等屬性值。"""
    parts = re.split(r"(<[^>]+>)", html_body)
    suit_map = {
        "♠": '<span class="s b">♠</span>',
        "♣": '<span class="s b">♣</span>',
        "♥": '<span class="s r">♥</span>',
        "♦": '<span class="s r">♦</span>',
    }
    for i, part in enumerate(parts):
        if part.startswith("<"):
            continue
        for suit, replacement in suit_map.items():
            part = part.replace(suit, replacement)
        parts[i] = part
    return "".join(parts)


BOOK_DOCS = [
    dict(
        md_file="3基本準則/第三章.md",
        slug="第三章",
        title="第三章 北歐梅花的問叫和基本準則",
        image_dir="第三章",
        lede="接力系統的正式定義：五種問叫工具（Sidestep／Splinter／54Pick-up／Six-shooter／CRASH）與四條硬規則的原始出處。",
    ),
    dict(
        md_file="4一方塊開叫/一方塊開叫.md",
        slug="一方塊開叫",
        title="第四章 一方塊開叫及其發展",
        image_dir="一方塊開叫",
        lede="1♦ 開叫的完整定義與 1♦–1♥ 後的接力叫牌發展，含原始截圖對照還原。",
    ),
    dict(
        md_file="5高花開叫/高花開叫.md",
        slug="高花開叫",
        title="第五章 一階高花開叫及其發展",
        image_dir="高花開叫",
        lede="1♠／1♥ 開叫後的第一應叫與 1NT 後的接力主幹發展，含原始截圖對照還原。",
    ),
    dict(
        md_file="7一梅花開叫/一梅花開叫.md",
        slug="一梅花開叫",
        title="第七章 一梅花開叫及其發展",
        image_dir="一梅花開叫",
        lede="1♣ 開叫後的應叫系統與各分支接力發展，含原始截圖對照還原。",
    ),
]


def build_book_pages():
    toc = "".join(
        '<a href="%s.html">%s</a>' % (d["slug"], d["title"].split(" ", 1)[-1] if " " in d["title"] else d["title"])
        for d in BOOK_DOCS
    )
    for d in BOOK_DOCS:
        md_path = os.path.join(SRC, d["md_file"])
        content = md_to_html(md_path, "../assets/images/%s/" % d["image_dir"])

        body = page_head(
            "原始書籍檔案", d["title"],
            ["Markdown 還原版", "來源：%s" % d["md_file"]],
            d["lede"],
            crumb_html='<div class="crumb"><a href="../index.html">首頁</a> ／ 原始書籍檔案 ／ %s</div>' % d["title"],
        )
        body += '<div class="toc-links">%s</div>' % toc
        body += '<article class="article">%s</article>' % content
        body += '<footer class="site-footer">整理自《北歐接力精確體系》原始剪藏，已比對截圖還原花色符號與表格結構。</footer>'

        html_out = render_page(
            title="%s ｜ 原始書籍檔案 ｜ 北歐接力精確體系" % d["title"],
            description=d["lede"],
            root_prefix="../",
            current_href="原始書籍檔案/%s.html" % d["slug"],
            body_html=body,
        )
        write_file(os.path.join(OUT, "原始書籍檔案", "%s.html" % d["slug"]), html_out)


# --------------------------------------------------------------------------
# 首頁
# --------------------------------------------------------------------------

def build_index():
    body = page_head(
        "北歐接力精確體系",
        "接力叫序資料庫",
        ["3 份叫序資料表", "4 篇原始書籍章節"],
        "整理自《北歐接力精確體系》一書，收錄一方塊、一梅花、高花三種開叫的完整接力叫序資料，"
        "並保留原始書籍章節的 Markdown 還原版供對照查閱。",
    )

    body += """
<section>
  <div class="sec-head"><h2>叫序資料</h2><span class="tag">依開叫花色整理的接力叫序表</span></div>
  <div class="card-grid">
    <a class="doc-card" href="叫序資料/一方塊開叫.html">
      <div class="k">Diamonds</div>
      <h3>一方塊開叫</h3>
      <p>1♦ 開叫、1♦–1♥ 接力主幹與各分支發展，共 12 個工作表。</p>
    </a>
    <a class="doc-card" href="叫序資料/一梅花開叫.html">
      <div class="k">Clubs</div>
      <h3>一梅花開叫</h3>
      <p>1♣ 開叫的應叫系統與各分支接力發展，共 13 個工作表。</p>
    </a>
    <a class="doc-card" href="叫序資料/高花開叫.html">
      <div class="k">Majors</div>
      <h3>高花開叫</h3>
      <p>一階高花開叫、1NT 後接力主幹與各分支，共 11 個工作表。</p>
    </a>
  </div>
</section>

<section>
  <div class="sec-head"><h2>原始書籍檔案</h2><span class="tag">Markdown 還原版，含原書截圖對照</span></div>
  <div class="card-grid">
    <a class="doc-card" href="原始書籍檔案/第三章.html">
      <div class="k">Chapter 03</div>
      <h3>第三章：問叫與基本準則</h3>
      <p>接力系統的正式定義：五種問叫工具與四條硬規則。</p>
    </a>
    <a class="doc-card" href="原始書籍檔案/一方塊開叫.html">
      <div class="k">Chapter 04</div>
      <h3>一方塊開叫及其發展</h3>
      <p>1♦ 開叫的完整定義與接力叫牌發展。</p>
    </a>
    <a class="doc-card" href="原始書籍檔案/高花開叫.html">
      <div class="k">Chapter 05</div>
      <h3>一階高花開叫及其發展</h3>
      <p>高花開叫的第一應叫與接力主幹發展。</p>
    </a>
    <a class="doc-card" href="原始書籍檔案/一梅花開叫.html">
      <div class="k">Chapter 07</div>
      <h3>一梅花開叫及其發展</h3>
      <p>1♣ 開叫後的應叫系統與各分支接力發展。</p>
    </a>
  </div>
</section>

<div class="external-card">
  <p><strong>接力工具速查</strong> — 第三章五種問叫工具（Sidestep／Splinter／54Pick-up／Six-shooter／CRASH）的濃縮速查頁。</p>
  <a href="北歐接力工具速查.html">開啟速查頁</a>
</div>
"""
    body += '<footer class="site-footer">整理自《北歐接力精確體系》。以上資料僅供個人叫牌系統學習與對照使用。</footer>'

    html_out = render_page(
        title="北歐接力精確體系 ｜ 接力叫序資料庫",
        description="整理自《北歐接力精確體系》的接力叫序資料庫與原始書籍檔案",
        root_prefix="",
        current_href=None,
        body_html=body,
    )
    write_file(os.path.join(OUT, "index.html"), html_out)


# --------------------------------------------------------------------------

def copy_quickref():
    src = os.path.join(SRC, "北歐接力工具速查.html")
    dst = os.path.join(OUT, "北歐接力工具速查.html")
    shutil.copyfile(src, dst)
    print("copied", os.path.relpath(dst, ROOT))


def main():
    copy_quickref()
    build_index()
    build_xlsx_page("一方塊開叫.xlsx", "一方塊開叫", "叫序資料 · 一方塊開叫",
                     "1♦ 開叫、1♦–1♥ 接力主幹（A–J 分支）與非接力叫牌、示例牌張、示例叫序。")
    build_xlsx_page("一梅花開叫.xlsx", "一梅花開叫", "叫序資料 · 一梅花開叫",
                     "1♣ 開叫後的應叫系統、1♦ 後開叫者再叫、1NT／2♣／2♦／2♥／2♠／2NT 各系統。")
    build_xlsx_page("高花開叫.xlsx", "高花開叫", "叫序資料 · 高花開叫",
                     "一階高花開叫的第一應叫、1NT 後接力主幹（A–E 分支）與示例牌張、示例叫序。")
    build_book_pages()


if __name__ == "__main__":
    main()
